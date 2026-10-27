"""Normalize Hande news workbooks; run with --input-dir PATH [--output-dir PATH]."""
import argparse
import csv
import re
import unicodedata
from collections import Counter
from pathlib import Path

HEADERS = ['title', 'year', 'type', 'scientific area', 'type of member', 'lab']
FIELDS = HEADERS + ['type_original', 'type_note', 'scientific_area_note',
                    'member_type_note', 'lab_note', 'year_note', 'source_file',
                    'row', 'slug', 'excerpt_ok', 'title_len']
SOURCES = [('00f1ff7a-UNIDCOM_News.xlsx', 'FCT Activities 2026'),
           ('71f47e3a-UNIDCOM_News_2025.xlsx', 'FCT Report 2026')]


def clean(value):
    text = '' if value is None else str(value).strip()
    return '' if text == '-' else text


def match_taxonomy(value, allowed):
    value = clean(value)
    matches = {item.casefold(): item for item in allowed}
    if not value:
        return '', ''
    return (matches[value.casefold()], '') if value.casefold() in matches else (value, 'non-matching taxonomy value')


def normalize_type(value, allowed):
    original = clean(value)
    canonical, note = match_taxonomy(original, allowed)
    if original and not note:
        return canonical, original, f'{original} -> {canonical}' if original != canonical else ''
    if original and not original.startswith('['):
        matches = [item for item in allowed
                   if re.sub(r'^\[[^]]+\]\s*', '', item).casefold() == original.casefold()]
        if len(matches) == 1:
            return matches[0], original, f'{original} -> {matches[0]}'
    return 'Unclassified', original, 'empty type' if not original else 'unknown type'


def normalize_year(value):
    value = clean(value)
    if value.casefold() == 'earlier':
        return 'Earlier', ''
    if re.fullmatch(r'[1-9][0-9]{3}(?:\.0)?', value):
        return value[:4], ''
    return '', '' if not value else f'invalid year: {value}'


def unique_slug(title, used):
    base = unicodedata.normalize('NFKD', title).encode('ascii', 'ignore').decode().lower()
    base = re.sub(r'[^a-z0-9]+', '-', base).strip('-') or 'news'
    slug, suffix = base, 2
    while slug in used:
        slug = f'{base}-{suffix}'
        suffix += 1
    used.add(slug)
    return slug


def normalize_rows(records):
    """Records contain source filename, Excel row number, six values, taxonomy."""
    normalized, review, seen, used = [], [], {}, set()
    for source, row_number, values, taxonomy in records:
        item = dict(zip(HEADERS, map(clean, values)))
        item['type'], item['type_original'], item['type_note'] = normalize_type(values[2], taxonomy['Type values'])
        item['year'], item['year_note'] = normalize_year(values[1])
        reasons = []
        for field, tax, note_field in [('scientific area', 'Scientific areas', 'scientific_area_note'),
                                       ('type of member', 'Member types', 'member_type_note'),
                                       ('lab', 'Labs', 'lab_note')]:
            item[field], item[note_field] = match_taxonomy(item[field], taxonomy[tax])
            if item[note_field]:
                reasons.append(f'{field}: {item[note_field]} ({item[field]})')
        if item['type'] == 'Unclassified':
            reasons.append(f'Unclassified type: {item["type_note"]}')
        if not item['lab']:
            reasons.append('missing lab')
        if item['year_note']:
            reasons.append(item['year_note'])
        if not item['title']:
            reasons.append('missing title')
        item.update(source_file=source, row=row_number,
                    slug=unique_slug(item['title'], used),
                    excerpt_ok=len(item['title']) <= 150, title_len=len(item['title']))
        if not item['excerpt_ok']:
            reasons.append('title > 150 chars')
        key = (item['title'], item['year'])
        if key in seen:
            reasons.append(f'duplicate title+year; retained {seen[key]}')
        else:
            seen[key] = f'{source}:row {row_number}'
            normalized.append(item)
        if reasons:
            review.append({**item, 'reason': '; '.join(reasons)})
    return normalized, review


def read_records(input_dir):
    from openpyxl import load_workbook
    for filename, sheet in SOURCES:
        workbook = load_workbook(input_dir / filename, read_only=True, data_only=True)
        try:
            tax_rows = list(workbook['Taxonomia'].iter_rows(values_only=True))
            header = next(i for i, row in enumerate(tax_rows) if clean(row[0]) == 'Year values')
            taxonomy = {clean(label): {clean(row[col]) for row in tax_rows[header + 1:] if clean(row[col])}
                        for col, label in enumerate(tax_rows[header]) if clean(label)}
            rows = workbook[sheet].iter_rows(values_only=True)
            assert [clean(v) for v in next(rows)] == HEADERS, f'Unexpected columns: {filename}'
            for number, values in enumerate(rows, 2):
                if any(clean(value) for value in values):
                    yield filename, number, values, taxonomy
        finally:
            workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', required=True, type=Path)
    parser.add_argument('--output-dir', type=Path, default=Path.home() / 'Backups/unidcom')
    args = parser.parse_args()
    records = list(read_records(args.input_dir))
    normalized, review = normalize_rows(records)
    out = args.output_dir
    out.mkdir(parents=True, exist_ok=True)
    for name, rows, fields in [('news_normalized.csv', normalized, FIELDS),
                               ('news_review.csv', review, FIELDS + ['reason'])]:
        with (out / name).open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
    print(f'Input: {len(records)}; normalized: {len(normalized)}; review: {len(review)}')
    for source in SOURCES:
        name = source[0]
        print(f'Source {name}: input={sum(r[0] == name for r in records)}, retained={sum(r["source_file"] == name for r in normalized)}')
    for field in ['type', 'lab']:
        print(f'Counts per {field} (retained rows):')
        for value, count in sorted(Counter(r[field] for r in normalized).items()):
            print(f'  {value or "(missing)"}: {count}')
    print(f'Unclassified: {sum(r["type"] == "Unclassified" for r in normalized)}')
    print(f'Missing lab: {sum(not r["lab"] for r in normalized)}')
    print(f'Duplicates dropped: {len(records) - len(normalized)}')
    print(f'Type mappings: {sum(" -> " in r["type_note"] for r in normalized)}')
    print(f'Taxonomy non-matches: {sum(any(r[k] for k in ["scientific_area_note", "member_type_note", "lab_note"]) for r in normalized)}')
    print(f'Titles >150: {sum(not r["excerpt_ok"] for r in normalized)}')


if __name__ == '__main__':
    main()
